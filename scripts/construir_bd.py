"""
Construye una base de datos relacional de aleaciones de aluminio a partir de
'Aleaciones_de_Al.xlsx'. Genera:
  - aleaciones_al.db   (SQLite)
  - aleaciones_al_BD.xlsx (mismas tablas en hojas + observaciones + vista)
Toda corrección o dato dudoso queda registrado en la tabla 'observacion'.
"""
import re, sqlite3, sys
from pathlib import Path
import openpyxl
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

SRC = Path(sys.argv[1])
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else Path(".")
wb = openpyxl.load_workbook(SRC)

# ---------------------------------------------------------------- fuentes
fuentes = pd.DataFrame([
    (1, "alu-stock.es – El aluminio", "https://www.alu-stock.es/es/informacion-tecnica/el-aluminio/"),
    (2, "AZoM – Aluminium alloys", "https://www.azom.com/article.aspx?ArticleID=2863"),
    (3, "alu-stock.es – Aleaciones", "https://www.alu-stock.es/es/informacion-tecnica/aleaciones/"),
], columns=["fuente_id", "nombre", "url"])

obs = []  # observaciones: (tabla, registro, campo, valor_original, accion, detalle)
def nota(tabla, registro, campo, original, accion, detalle):
    obs.append(dict(tabla=tabla, registro=str(registro), campo=campo,
                    valor_original=None if original is None else str(original),
                    accion=accion, detalle=detalle))

def num(s):
    """'0,05' -> 0.05 ; admite números ya numéricos."""
    if s is None: return None
    if isinstance(s, (int, float)): return float(s)
    s = str(s).strip().replace("\xa0", "").replace(",", ".")
    try: return float(s)
    except ValueError: return None

# ---------------------------------------------------------------- series
series = pd.DataFrame([
    ("1XXX", "Ninguno (Al ≥ 99,00 %)", "No"),
    ("2XXX", "Cobre (Cu)", "Sí"),
    ("3XXX", "Manganeso (Mn)", "No"),
    ("4XXX", "Silicio (Si)", "Variable"),
    ("5XXX", "Magnesio (Mg)", "No"),
    ("6XXX", "Magnesio + Silicio (Mg2Si)", "Sí"),
    ("7XXX", "Zinc (Zn)", "Sí"),
    ("8XXX", "Otros elementos (Li, Fe, Sn…)", "Variable"),
], columns=["serie", "elemento_principal", "tratable_termicamente"])
nota("serie", "8XXX", "elemento_principal", "Litio", "corregido",
     "La serie 8XXX agrupa 'otros elementos'; las Al-Li (p.ej. 8090) son solo una parte.")
nota("serie", "*", "tratable_termicamente", None, "agregado",
     "Columna agregada (conocimiento general de metalurgia), no estaba en el original.")

# ---------------------------------------------------------------- estados
ws = wb["Clasificación"]
est_rows = []
seccion = None
for r in range(19, 83):
    v = ws.cell(r, 2).value
    if not v: continue
    v = v.replace("\xa0", " ")
    if "|" not in v:
        continue
    cab, desc = [x.strip() for x in v.split("|", 1)]
    desc = desc.replace("a partir de 0 ó F", "a partir de O ó F")
    if ":" in cab:
        cod, nombre = [x.strip() for x in cab.split(":", 1)]
    else:
        cod, nombre = cab, ""
    cod = cod.replace(" ", "")
    for c in [c.strip() for c in cod.split(",")]:
        est_rows.append(dict(codigo=c, tipo=c[0], nombre=nombre or None,
                             descripcion=desc, origen="Hoja Clasificación"))
estados = pd.DataFrame(est_rows)
# T53 aparece en el grupo T51–T54 y luego con definición propia → queda la específica
dup = estados[estados.codigo.duplicated(keep=False)]
if len(dup):
    nota("estado", "T53", "descripcion", None, "resuelto",
         "T53 aparece dos veces en el original (grupo T51–T54 y definición propia); se conservó la propia.")
    estados = estados.drop_duplicates("codigo", keep="last")
nota("estado", "T42/T62", "descripcion", "a partir de 0 ó F", "corregido", "'0' era la letra O (recocido).")
# Estados "genéricos" que usan las tablas de propiedades
genericos = [
    ("H (todos)", "H", "Todos los estados H", "Valor válido para cualquier estado H (así figura en la fuente)."),
    ("H (resto)", "H", "Resto de estados H", "Estados H distintos de los listados por separado para esa aleación."),
    ("T (todos)", "T", "Todos los estados T", "Valor válido para cualquier estado T (así figura en la fuente)."),
]
estados = pd.concat([estados, pd.DataFrame(
    [dict(codigo=c, tipo=t, nombre=n, descripcion=d, origen="Agregado (agrupa estados)") for c, t, n, d in genericos])],
    ignore_index=True)
estados = estados.reset_index(drop=True)

# ---------------------------------------------------------------- composición
ws = wb["Comp Quím"]
cols = [ws.cell(6, c).value for c in range(3, 22)]  # C..U
ren = {"Otro": "Otros (c/u)", "Total": "Otros (total)"}
comp_rows, aleaciones = [], {}
for r in range(8, ws.max_row + 1):
    a = ws.cell(r, 2).value
    if a is None: continue
    a = str(int(a))
    aleaciones[a] = True
    for i, el in enumerate(cols):
        raw = ws.cell(r, 3 + i).value
        if raw is None: continue
        s = str(raw).strip().replace("\xa0", "")
        if s in ("-", ""): continue
        el = ren.get(el, el)
        mn = mx = None; balance = 0
        if s.lower() == "resto":
            balance = 1
        elif s.startswith("≤"):
            mx = num(s[1:])
        elif s.startswith("≥"):
            mn = num(s[1:])
        elif "-" in s:
            p = s.split("-"); mn, mx = num(p[0]), num(p[1])
        else:
            mn = mx = num(s)
        if a == "1050" and el == "Al" and s.startswith("≤"):
            nota("composicion", "1050", "Al", s, "corregido",
                 "Al ≤99,50 es un error: la 1050 exige Al ≥ 99,50 %.")
            mn, mx = 99.50, None
        comp_rows.append(dict(aleacion=a, elemento=el, pct_min=mn, pct_max=mx,
                              es_balance=balance, valor_original=s, fuente_id=3))
composicion = pd.DataFrame(comp_rows)

# ---------------------------------------------------------------- físicas
ws = wb["Prop Físicas"]
map_estado = {"Todos H": ["H (todos)"], "Resto H": ["H (resto)"], "Todos T": ["T (todos)"]}
fis_rows = []
for r in range(8, ws.max_row + 1):
    a = ws.cell(r, 2).value
    if a is None: continue
    a = str(int(a)); e_raw = str(ws.cell(r, 3).value).strip()
    ests = map_estado.get(e_raw) or [x.strip() for x in e_raw.split("/")]
    if len(ests) > 1:
        nota("prop_fisicas", f"{a} {e_raw}", "estado", e_raw, "desdoblado",
             f"Fila agrupada separada en {len(ests)} filas con los mismos valores.")
    fus = str(ws.cell(r, 7).value).split("-")
    v = lambda c: ws.cell(r, c).value
    for e in ests:
        fis_rows.append(dict(
            aleacion=a, estado=e,
            E_GPa=v(4) / 1000 if v(4) else None, G_GPa=v(5) / 1000 if v(5) else None,
            poisson=v(6), T_solidus_C=num(fus[0]), T_liquidus_C=num(fus[1]),
            cp_J_kgK=v(8), alfa_um_mK=v(9), densidad_kg_m3=v(10),
            resistividad_nOhm_m=v(11), k_termica_W_mK=v(12), conductividad_IACS=v(13),
            fuente_id=3))
    if v(11) is None:
        nota("prop_fisicas", f"{a} {e_raw}", "resistividad / k / IACS", None, "faltante",
             "Sin datos de resistividad, conductividad térmica ni eléctrica.")
prop_fisicas = pd.DataFrame(fis_rows)
nota("prop_fisicas", "*", "E, G", "MPa", "convertido", "Módulos convertidos de MPa a GPa.")
nota("prop_fisicas", "*", "intervalo de fusión", "'645 - 658'", "separado",
     "Texto separado en T_solidus_C y T_liquidus_C.")

# ---------------------------------------------------------------- mecánicas
ws = wb["Prop Mecánicas"]
mec_rows, actual = [], None
for r in range(9, ws.max_row + 1):
    if ws.cell(r, 2).value is not None:
        actual = str(int(ws.cell(r, 2).value))
    t = ws.cell(r, 3).value
    if t is None: continue
    t_orig = t
    if t in (0, 0.0, "0"):
        t = "O"
    t = str(t).strip()
    if re.fullmatch(r"H\d", t):
        t = "HX" + t[1]
    v = lambda c: ws.cell(r, c).value
    mec_rows.append(dict(aleacion=actual, estado=t, Rp02_MPa=v(6), Rm_MPa=v(7), tau_MPa=v(8),
                         A5_pct=v(9), A50_pct=v(10), HB=v(11), HV=v(12), fatiga_MPa=v(13),
                         dato_dudoso=0, fuente_id=2))
prop_mecanicas = pd.DataFrame(mec_rows)
nota("prop_mecanicas", "*", "estado", "0", "corregido",
     "El estado recocido 'O' estaba cargado como número 0.")
nota("prop_mecanicas", "*", "estado", "H2…H9", "renombrado",
     "Se registró como HX2…HX9: la fuente no indica si es H1x, H2x o H3x.")
nota("prop_mecanicas", "*", "aleacion", None, "completado",
     "El número de aleación solo figuraba en la primera fila de cada grupo; se completó hacia abajo.")
nota("prop_mecanicas", "*", "Composición química", None, "eliminado",
     "Columnas D:E vacías en el original; la composición está en la tabla 'composicion'.")
for a, e, campo, det in [
    ("5251", "O", "Rp02_MPa / Rm_MPa", "Rp0,2=145 y Rm=300 MPa en estado recocido superan a los estados H2–H8: probable error de carga (esperable Rm ≈ 170–200 MPa)."),
    ("6063", "O", "HV", "HV=85 no es coherente con HB=25 (deberían ser similares)."),
]:
    m = (prop_mecanicas.aleacion == a) & (prop_mecanicas.estado == e)
    prop_mecanicas.loc[m, "dato_dudoso"] = 1
    nota("prop_mecanicas", f"{a}-{e}", campo, None, "marcado dudoso", det + " No se modificó el valor; verificar en otra fuente.")

# ---------------------------------------------------------------- aleaciones (maestra)
todas = sorted(set(aleaciones) | set(prop_mecanicas.aleacion))
aleacion = pd.DataFrame([dict(aleacion=a, serie=a[0] + "XXX",
                              tiene_composicion=int(a in set(composicion.aleacion)),
                              tiene_prop_fisicas=int(a in set(prop_fisicas.aleacion)),
                              tiene_prop_mecanicas=int(a in set(prop_mecanicas.aleacion)))
                         for a in todas])
for a in aleacion.query("tiene_composicion==0").aleacion:
    nota("aleacion", a, "*", None, "faltante", "Solo tiene propiedades mecánicas; faltan composición y propiedades físicas.")
sin_mec = aleacion.query("tiene_prop_mecanicas==0").aleacion.tolist()
nota("aleacion", ", ".join(sin_mec), "prop_mecanicas", None, "faltante",
     f"{len(sin_mec)} aleaciones sin propiedades mecánicas.")

# ---------------------------------------------------------------- Al puro
ws = wb["Prop Generales"]
puro = []
for r in range(7, 42):
    p, val = ws.cell(r, 2).value, ws.cell(r, 3).value
    if p: puro.append(dict(propiedad=str(p).strip(), valor=str(val).strip(), fuente_id=1))
for r in range(6, 18):
    p, val = ws.cell(r, 6).value, ws.cell(r, 10).value
    if p: puro.append(dict(propiedad=str(p).strip(), valor=str(val).strip(), fuente_id=2))
prop_al_puro = pd.DataFrame(puro)
prop_al_puro["valor_num"] = prop_al_puro.valor.map(num)
fix = {  # propiedad (fuente) : (nuevo valor, detalle)
    ("Densidad (g/cm3)", 2): ("2.698", "26,898 era un error de tipeo (densidad del Al = 2,698 g/cm³)."),
    ("Resistividad elécrica a 20 C (Ohm/cm)", 2): ("2.69", "La unidad correcta es µΩ·cm, no Ohm/cm."),
    ("Volumen molar", 1): ("9,99 cm3/mol", "'cm3/topo' es una mala traducción de 'mol'."),
    ("Densidad", 1): ("2.7", "Unidad faltante: g/cm³."),
}
for (p, f), (nv, det) in fix.items():
    m = (prop_al_puro.propiedad == p) & (prop_al_puro.fuente_id == f)
    nota("prop_al_puro", p, "valor", prop_al_puro.loc[m, "valor"].iloc[0], "corregido", det)
    prop_al_puro.loc[m, "valor"] = nv
    prop_al_puro.loc[m, "valor_num"] = num(nv.split()[0])
ren_p = {"Resistividad elécrica a 20 C (Ohm/cm)": "Resistividad eléctrica a 20 °C (µΩ·cm)",
         "Densidad": "Densidad (g/cm3)"}
prop_al_puro["propiedad"] = prop_al_puro.propiedad.replace(ren_p)
nota("prop_al_puro", "*", "*", None, "marcado",
     "Las dos fuentes difieren en: módulo elástico (68,3 vs 70 GPa), ebullición (2467 vs 2480 °C), "
     "electronegatividad (1,5 vs 1,61). Se conservan ambos valores con su fuente.")

# ---------------------------------------------------------------- chequeos de integridad
assert set(prop_fisicas.estado) <= set(estados.codigo), set(prop_fisicas.estado) - set(estados.codigo)
assert set(prop_mecanicas.estado) <= set(estados.codigo), set(prop_mecanicas.estado) - set(estados.codigo)
observacion = pd.DataFrame(obs)
observacion.insert(0, "obs_id", range(1, len(observacion) + 1))

# ---------------------------------------------------------------- SQLite
db = OUT / "aleaciones_al.db"
db.unlink(missing_ok=True)
con = sqlite3.connect(db)
con.executescript("""
PRAGMA foreign_keys = ON;
CREATE TABLE fuente (fuente_id INTEGER PRIMARY KEY, nombre TEXT, url TEXT);
CREATE TABLE serie (serie TEXT PRIMARY KEY, elemento_principal TEXT, tratable_termicamente TEXT);
CREATE TABLE aleacion (aleacion TEXT PRIMARY KEY, serie TEXT REFERENCES serie(serie),
  tiene_composicion INT, tiene_prop_fisicas INT, tiene_prop_mecanicas INT);
CREATE TABLE estado (codigo TEXT PRIMARY KEY, tipo TEXT, nombre TEXT, descripcion TEXT, origen TEXT);
CREATE TABLE composicion (aleacion TEXT REFERENCES aleacion(aleacion), elemento TEXT,
  pct_min REAL, pct_max REAL, es_balance INT, valor_original TEXT,
  fuente_id INT REFERENCES fuente(fuente_id), PRIMARY KEY (aleacion, elemento));
CREATE TABLE prop_fisicas (aleacion TEXT REFERENCES aleacion(aleacion), estado TEXT REFERENCES estado(codigo),
  E_GPa REAL, G_GPa REAL, poisson REAL, T_solidus_C REAL, T_liquidus_C REAL, cp_J_kgK REAL,
  alfa_um_mK REAL, densidad_kg_m3 REAL, resistividad_nOhm_m REAL, k_termica_W_mK REAL,
  conductividad_IACS REAL, fuente_id INT REFERENCES fuente(fuente_id), PRIMARY KEY (aleacion, estado));
CREATE TABLE prop_mecanicas (aleacion TEXT REFERENCES aleacion(aleacion), estado TEXT REFERENCES estado(codigo),
  Rp02_MPa REAL, Rm_MPa REAL, tau_MPa REAL, A5_pct REAL, A50_pct REAL, HB REAL, HV REAL,
  fatiga_MPa REAL, dato_dudoso INT, fuente_id INT REFERENCES fuente(fuente_id), PRIMARY KEY (aleacion, estado));
CREATE TABLE prop_al_puro (propiedad TEXT, valor TEXT, valor_num REAL, fuente_id INT REFERENCES fuente(fuente_id));
CREATE TABLE observacion (obs_id INTEGER PRIMARY KEY, tabla TEXT, registro TEXT, campo TEXT,
  valor_original TEXT, accion TEXT, detalle TEXT);
""")
tablas = dict(fuente=fuentes, serie=series, aleacion=aleacion, estado=estados, composicion=composicion,
              prop_fisicas=prop_fisicas, prop_mecanicas=prop_mecanicas, prop_al_puro=prop_al_puro,
              observacion=observacion)
for n, df in tablas.items():
    df.to_sql(n, con, if_exists="append", index=False)
# Vista: aleación + estado con físicas y mecánicas juntas
con.executescript("""
CREATE VIEW v_propiedades AS
SELECT a.aleacion, a.serie, s.elemento_principal, m.estado,
       f.densidad_kg_m3, f.E_GPa, f.conductividad_IACS, f.k_termica_W_mK,
       m.Rp02_MPa, m.Rm_MPa, m.A50_pct, m.HB, m.fatiga_MPa, m.dato_dudoso
FROM prop_mecanicas m
JOIN aleacion a ON a.aleacion = m.aleacion
JOIN serie s ON s.serie = a.serie
LEFT JOIN prop_fisicas f ON f.aleacion = m.aleacion
  AND (f.estado = m.estado
       OR (f.estado IN ('H (todos)', 'H (resto)') AND m.estado LIKE 'H%')
       OR (f.estado = 'T (todos)' AND m.estado LIKE 'T%'))
UNION ALL
SELECT a.aleacion, a.serie, s.elemento_principal, f.estado,
       f.densidad_kg_m3, f.E_GPa, f.conductividad_IACS, f.k_termica_W_mK,
       NULL, NULL, NULL, NULL, NULL, NULL
FROM prop_fisicas f
JOIN aleacion a ON a.aleacion = f.aleacion
JOIN serie s ON s.serie = a.serie
WHERE f.aleacion NOT IN (SELECT aleacion FROM prop_mecanicas);
""")
fk = con.execute("PRAGMA foreign_key_check").fetchall()
assert not fk, fk
vista = pd.read_sql("SELECT * FROM v_propiedades ORDER BY aleacion, estado", con)
con.commit(); con.close()

# ---------------------------------------------------------------- Excel
xlsx = OUT / "aleaciones_al_BD.xlsx"
leeme = pd.DataFrame({"Hoja": [
    "aleacion", "serie", "estado", "composicion", "prop_fisicas", "prop_mecanicas",
    "prop_al_puro", "fuente", "observacion", "vista_propiedades"],
    "Contenido": [
    "Tabla maestra: una fila por aleación, con indicadores de qué datos tiene.",
    "Familias 1XXX–8XXX y su elemento principal.",
    "Diccionario de estados/temples (F, O, H, T y subdivisiones).",
    "Composición química en formato largo: una fila por aleación y elemento (mín/máx en %).",
    "Propiedades físicas por aleación y estado.",
    "Propiedades mecánicas por aleación y estado. dato_dudoso=1 → revisar.",
    "Propiedades del aluminio puro, de ambas fuentes.",
    "Fuentes bibliográficas (fuente_id se usa en las demás tablas).",
    "Registro de todas las correcciones, datos dudosos y faltantes.",
    "Vista combinada (solo lectura) de físicas + mecánicas para filtrar y comparar."],
    "Clave / relación": [
    "aleacion (PK)", "serie (PK)", "codigo (PK)", "aleacion + elemento",
    "aleacion + estado → aleacion, estado", "aleacion + estado → aleacion, estado",
    "fuente_id → fuente", "fuente_id (PK)", "obs_id (PK)", "—"]})
hojas = {"LEEME": leeme, **{k: tablas[k] for k in
         ["aleacion", "serie", "estado", "composicion", "prop_fisicas", "prop_mecanicas",
          "prop_al_puro", "fuente", "observacion"]}, "vista_propiedades": vista}
with pd.ExcelWriter(xlsx, engine="openpyxl") as w:
    for n, df in hojas.items():
        df.to_excel(w, sheet_name=n, index=False)
wbo = openpyxl.load_workbook(xlsx)
hdr_fill = PatternFill("solid", fgColor="1F4E78")
dudoso_fill = PatternFill("solid", fgColor="FFF2CC")
for wsx in wbo.worksheets:
    for row in wsx.iter_rows():
        for c in row:
            c.font = Font(name="Arial", size=10, bold=c.row == 1, color="FFFFFF" if c.row == 1 else "000000")
            if c.row == 1: c.fill = hdr_fill
    wsx.freeze_panes = "A2"
    if wsx.title != "LEEME": wsx.auto_filter.ref = wsx.dimensions
    for i, col in enumerate(wsx.columns, 1):
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        wsx.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 70)
        if width > 70:
            for c in col: c.alignment = Alignment(wrap_text=True, vertical="top")
    hdrs = [c.value for c in wsx[1]]
    if "dato_dudoso" in hdrs:
        k = hdrs.index("dato_dudoso")
        for row in wsx.iter_rows(min_row=2):
            if row[k].value == 1:
                for c in row: c.fill = dudoso_fill
wbo.save(xlsx)
print({n: len(df) for n, df in tablas.items()}, "vista:", len(vista))
